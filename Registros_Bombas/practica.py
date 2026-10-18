import serial
import time
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
import os
import re

# ================= CONFIGURACIÓN =================
# Puerto USB configurado 
PUERTO_COM = '/dev/ttyUSB0' 
BAUD_RATE = 115200
ARCHIVO_EXCEL = "Registro_Paros_Bombas.xlsx"
# =================================================

# Diccionario para recordar el estado de cada bomba y no repetir registros
estado_actual = {1: "OFF", 2: "OFF", 3: "OFF", 4: "OFF"}

def inicializar_excel():
    """Crea el archivo Excel con encabezados si no existe"""
    if not os.path.exists(ARCHIVO_EXCEL):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial de Paros"
        
        # Encabezados
        ws.append(["Fecha", "Hora", "Bomba", "Evento", "Corriente Registrada"])
        
        # Dar formato a los encabezados
        bold_font = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="4F81BD")
        for cell in ws["1:1"]:
            cell.font = bold_font
            cell.fill = fill
            
        wb.save(ARCHIVO_EXCEL)
        print(f"[*] Archivo {ARCHIVO_EXCEL} creado exitosamente.")

def registrar_evento(bomba, evento, valor):
    """Guarda una nueva fila en el Excel con la fecha y hora exacta"""
    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    ws = wb.active
    
    ahora = datetime.now()
    fecha = ahora.strftime("%d/%m/%Y")
    hora = ahora.strftime("%H:%M:%S")
    
    ws.append([fecha, hora, f"Bomba {bomba}", evento, valor])
    wb.save(ARCHIVO_EXCEL)
    print(f"[EXCEL] Guardado: {fecha} {hora} | Bomba {bomba} | {evento} ({valor})")

def iniciar_monitoreo():
    inicializar_excel()
    
    try:
        ser = serial.Serial(PUERTO_COM, BAUD_RATE, timeout=1)
        print(f"\n[+] Conectado al puerto {PUERTO_COM}. Escuchando datos...")
        print("Presiona Ctrl+C para detener el registro.\n")
        time.sleep(2) # Esperar a que reinicie el ESP32
        
        while True:
            if ser.in_waiting > 0:
                linea = ser.readline().decode('utf-8', errors='ignore').strip()
                
                # Buscamos los 4 valores en la línea de texto del ESP32
                valores = re.findall(r'POT \d: ([0-9]+|OFF)', linea)
                
                if len(valores) == 4:
                    for i, val_str in enumerate(valores):
                        bomba_id = i + 1
                        
                        if val_str == "OFF":
                            nuevo_estado = "OFF"
                        else:
                            corriente = int(val_str)
                            if corriente > 3000:
                                nuevo_estado = "PARO"
                            else:
                                nuevo_estado = "NORMAL"
                        
                        # Si el estado cambió, lo registramos
                        if nuevo_estado != estado_actual[bomba_id]:
                            if nuevo_estado == "PARO":
                                registrar_evento(bomba_id, " ALERTA: SOBRECORRIENTE", val_str)
                            elif nuevo_estado == "NORMAL" and estado_actual[bomba_id] == "PARO":
                                registrar_evento(bomba_id, " NORMALIZADO", val_str)
                            
                            # Actualizamos la memoria
                            estado_actual[bomba_id] = nuevo_estado

    except serial.SerialException as e:
        print(f"\n[!] ERROR de conexión con {PUERTO_COM}.")
        print("    -> Verifica que el Arduino IDE tenga el Monitor Serie CERRADO.")
        print("    -> Si el error dice 'Permission denied', abre otra terminal y ejecuta:")
        print(f"       sudo chmod a+rw {PUERTO_COM}")
    except KeyboardInterrupt:
        print("\n[!] Monitoreo detenido por el usuario.")
        if 'ser' in locals() and ser.is_open:
            ser.close()

if __name__ == "__main__":
    iniciar_monitoreo()
